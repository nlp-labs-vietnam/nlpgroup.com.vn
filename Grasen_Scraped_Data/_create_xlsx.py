"""Generate Grasen_News_Data.xlsx from scraped article metadata."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

articles = [
    {
        "ID": "POST_001",
        "Title": "EV Charging Management Software: Managing a Charging Network from One Backend",
        "Category": "News",
        "Published Date": "2026-09-30",
        "Author/Source": "GRASEN",
        "Source URL": "https://www.grasen.com/news/ev-charging-management-software",
        "Excerpt": "Learn how EV charging management software helps operators monitor chargers, manage users, pricing, payments and remote operations across multiple sites.",
        "Full Content": "See raw_articles/ev-charging-management-software.txt",
        "Tags": "EV charging management software; CSMS; charging network; OCPP; backend",
        "Thumbnail URL": "https://upload.grasen.com/upload/2026/09/30/EV Charging Management Software_20260930103919A129.png",
        "Thumbnail Local Path": "images/ev-charging-management-software/cover.png",
        "Content Images": "(none — JS-rendered page)",
    },
    {
        "ID": "POST_002",
        "Title": "Are All Electric Vehicle Chargers the Same? Not Quite—Here's Why",
        "Category": "News",
        "Published Date": "2026-09-28",
        "Author/Source": "GRASEN",
        "Source URL": "https://www.grasen.com/news/are-all-electric-vehicle-chargers-the-same",
        "Excerpt": "Learn how EV chargers differ in AC vs DC charging, power, connectors, vehicle compatibility and commercial charging features.",
        "Full Content": "See raw_articles/are-all-electric-vehicle-chargers-the-same.txt",
        "Tags": "EV charger types; AC charging; DC fast charging; connector standards; CCS; CHAdeMO; NACS",
        "Thumbnail URL": "https://upload.grasen.com/upload/2026/09/28/All Electric Vehicle Chargers the Same？1_20260928140207A128.jpg",
        "Thumbnail Local Path": "images/are-all-electric-vehicle-chargers-the-same/cover.jpg",
        "Content Images": "(none — JS-rendered page)",
    },
    {
        "ID": "POST_003",
        "Title": "Business Electric Car Charger Guide: Choosing the Right Solution for Your Site",
        "Category": "News",
        "Published Date": "2026-09-23",
        "Author/Source": "GRASEN",
        "Source URL": "https://www.grasen.com/news/business-electric-car-charger-guide",
        "Excerpt": "Learn how to choose the right EV charger for workplaces, hotels, fleets, retail sites and public charging projects.",
        "Full Content": "See raw_articles/business-electric-car-charger-guide.txt",
        "Tags": "business EV charger; commercial EV charging; workplace charging; fleet charging; hotel EV charger",
        "Thumbnail URL": "https://upload.grasen.com/upload/2026/09/24/T480000_20260924162451A126.png",
        "Thumbnail Local Path": "images/business-electric-car-charger-guide/cover.png",
        "Content Images": "(none — JS-rendered page)",
    },
    {
        "ID": "POST_004",
        "Title": "EV Charging Monetization: 6 Ways Commercial Charging Stations Generate Revenue",
        "Category": "News",
        "Published Date": "2026-09-23",
        "Author/Source": "GRASEN",
        "Source URL": "https://www.grasen.com/news/ev-charging-monetization",
        "Excerpt": "Explore six practical ways commercial EV charging stations can generate revenue through charging fees, smart pricing, advertising, memberships and software.",
        "Full Content": "See raw_articles/ev-charging-monetization.txt",
        "Tags": "EV charging monetization; revenue generation; CPO business model; charging fees; smart pricing; advertising screens",
        "Thumbnail URL": "https://upload.grasen.com/upload/2026/09/23/EV Charging Monetization picture_20260923095216A122.png",
        "Thumbnail Local Path": "images/ev-charging-monetization/cover.png",
        "Content Images": "(none — JS-rendered page)",
    },
    {
        "ID": "POST_005",
        "Title": "Top 10 DC Fast Charger Manufacturers in China (2026): Global Buyer's Guide",
        "Category": "News",
        "Published Date": "2026-09-22",
        "Author/Source": "GRASEN",
        "Source URL": "https://www.grasen.com/news/top-10-dc-fast-charger-manufacturers-china",
        "Excerpt": "Compare 10 DC fast charger manufacturers in China and see their main strengths, product focus and suitable project applications.",
        "Full Content": "See raw_articles/top-10-dc-fast-charger-manufacturers-china.txt",
        "Tags": "DC fast charger manufacturers; China EV charger; charging station suppliers; EVSE manufacturers; global buyer guide",
        "Thumbnail URL": "https://upload.grasen.com/upload/2026/09/22/Top 10 DC Fast Charger Manufacturer_20260922140219A121.png",
        "Thumbnail Local Path": "images/top-10-dc-fast-charger-manufacturers-china/cover.png",
        "Content Images": "(none — JS-rendered page)",
    },
    {
        "ID": "POST_006",
        "Title": "How to Use the Grasen App for EV Charging and Payment",
        "Category": "News",
        "Published Date": "2026-09-18",
        "Author/Source": "GRASEN",
        "Source URL": "https://www.grasen.com/news/how-to-use-grasen-app",
        "Excerpt": "Learn how to use the Grasen App to find nearby chargers, check charging status and prices, start a session, make payments and review your charging records.",
        "Full Content": "See raw_articles/how-to-use-grasen-app.txt",
        "Tags": "Grasen App; EV charging app; mobile payment; find charger; charging session management",
        "Thumbnail URL": "https://upload.grasen.com/upload/2026/09/18/Explore Grasen EV Charging Solutions_20260918114725A119.png",
        "Thumbnail Local Path": "images/how-to-use-grasen-app/cover.png",
        "Content Images": "(none — JS-rendered page)",
    },
    {
        "ID": "POST_007",
        "Title": "Grasen T480Q 480kW 4-Gun DC EV Fast Charger for High-Traffic Charging Sites",
        "Category": "News",
        "Published Date": "2026-09-16",
        "Author/Source": "GRASEN",
        "Source URL": "https://www.grasen.com/news/grasen-t480q-480kw-4-gun-dc-ev-fast-charger",
        "Excerpt": "Discover the new Grasen T480Q 480kW 4-gun DC fast charger, designed for high-traffic commercial charging sites with flexible connector configurations.",
        "Full Content": "See raw_articles/grasen-t480q-480kw-4-gun-dc-fast-charger.txt",
        "Tags": "T480Q; 480kW; 4-gun DC charger; high-traffic charging; multi-connector DC charger; Grasen product",
        "Thumbnail URL": "https://upload.grasen.com/upload/2026/09/23/480kW 4-Gun DC EV Fast Charger picture_20260923102825A123.png",
        "Thumbnail Local Path": "images/grasen-t480q-480kw-4-gun-dc-fast-charger/cover.png",
        "Content Images": "(none — JS-rendered page)",
    },
    {
        "ID": "POST_008",
        "Title": "Dual-Gun DC Fast Charger: How Does Power Sharing Work?",
        "Category": "News",
        "Published Date": "2026-09-15",
        "Author/Source": "GRASEN",
        "Source URL": "https://www.grasen.com/news/dual-gun-dc-fast-charger-power-sharing",
        "Excerpt": "How does a dual-gun DC fast charger divide power when two EVs charge at the same time? This guide explains equal and dynamic power sharing, vehicle demand, and what buyers should check before choosing a dual-connector charger.",
        "Full Content": "See raw_articles/dual-gun-dc-fast-charger-power-sharing.txt",
        "Tags": "dual-gun DC charger; power sharing; dynamic power distribution; EV charging; multi-gun charger",
        "Thumbnail URL": "https://upload.grasen.com/upload/2026/09/15/文章配图1)_20260915152149A115.jpg",
        "Thumbnail Local Path": "images/dual-gun-dc-fast-charger-power-sharing/cover.jpg",
        "Content Images": "(none — JS-rendered page)",
    },
    {
        "ID": "POST_009",
        "Title": "CCS2 DC Fast Chargers for Commercial Charging Stations: What Buyers Should Know",
        "Category": "News",
        "Published Date": "2026-09-10",
        "Author/Source": "GRASEN",
        "Source URL": "https://www.grasen.com/news/ccs2-dc-fast-chargers-commercial-stations",
        "Excerpt": "A guide to selecting CCS2 DC fast chargers for commercial charging stations, fleets and EV infrastructure projects.",
        "Full Content": "See raw_articles/ccs2-dc-fast-chargers-commercial-stations.txt",
        "Tags": "CCS2; DC fast charger; commercial charging stations; Combined Charging System; EV infrastructure; CPO",
        "Thumbnail URL": "https://upload.grasen.com/upload/2026/09/12/文章配图3_20260912085619A114.png",
        "Thumbnail Local Path": "images/ccs2-dc-fast-chargers-commercial-stations/cover.png",
        "Content Images": "(none — JS-rendered page)",
    },
    {
        "ID": "POST_010",
        "Title": "Invitation: Meet GRASEN at the 140th Canton Fair 2026",
        "Category": "Events",
        "Published Date": "2026-09-08",
        "Author/Source": "GRASEN",
        "Source URL": "https://www.grasen.com/news/grasen-canton-fair-140th-2026",
        "Excerpt": "GRASEN will exhibit at the 140th Canton Fair in Guangzhou from October 15-19, 2026. Visit Booth 8.1G20 to discover our new 480kW Four-Gun DC Fast Charger, R Series DC chargers, AC chargers, and EV charging solutions.",
        "Full Content": "See raw_articles/grasen-canton-fair-140th-2026.txt",
        "Tags": "Canton Fair 2026; trade show; exhibition; Guangzhou; GRASEN event; EV charger exhibition",
        "Thumbnail URL": "https://upload.grasen.com/upload/2026/09/10/广交会官网_20260910145319A112.jpg",
        "Thumbnail Local Path": "images/grasen-canton-fair-140th-2026/cover.jpg",
        "Content Images": "(none — JS-rendered page)",
    },
]

columns = [
    "ID", "Title", "Category", "Published Date", "Author/Source",
    "Source URL", "Excerpt", "Full Content", "Tags",
    "Thumbnail URL", "Thumbnail Local Path", "Content Images"
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Grasen News"

# Header style
header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF", size=11)
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

# Write headers
for col_idx, col_name in enumerate(columns, start=1):
    cell = ws.cell(row=1, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = thin_border

# Row styles
row_fills = [
    PatternFill(start_color="EEF2F7", end_color="EEF2F7", fill_type="solid"),
    PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"),
]

# Write data rows
for row_idx, article in enumerate(articles, start=2):
    fill = row_fills[(row_idx % 2)]
    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=article.get(col_name, ""))
        cell.fill = fill
        cell.border = thin_border
        cell.alignment = Alignment(vertical="top", wrap_text=True)

# Set column widths
col_widths = {
    "ID": 10,
    "Title": 60,
    "Category": 12,
    "Published Date": 16,
    "Author/Source": 14,
    "Source URL": 65,
    "Excerpt": 70,
    "Full Content": 40,
    "Tags": 60,
    "Thumbnail URL": 80,
    "Thumbnail Local Path": 55,
    "Content Images": 30,
}
for col_idx, col_name in enumerate(columns, start=1):
    ws.column_dimensions[get_column_letter(col_idx)].width = col_widths.get(col_name, 20)

# Freeze header row
ws.freeze_panes = "A2"

# Row heights for data rows
ws.row_dimensions[1].height = 30
for row_idx in range(2, len(articles) + 2):
    ws.row_dimensions[row_idx].height = 60

wb.save("Grasen_News_Data.xlsx")
print("Created: Grasen_News_Data.xlsx")
